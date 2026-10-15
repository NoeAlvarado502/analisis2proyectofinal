import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(element)
        else:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)

def format_apa_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = table.rows
    num_rows = len(rows)
    for r_idx, row in enumerate(rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.first_line_indent = Inches(0)
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if r_idx == 0:
                        run.font.bold = True
            if r_idx == 0:
                set_cell_border(cell, top={'val': 'single', 'sz': 8}, bottom={'val': 'single', 'sz': 6})
            elif r_idx == num_rows - 1:
                set_cell_border(cell, bottom={'val': 'single', 'sz': 8})
            else:
                set_cell_border(cell, top=None, bottom=None, left=None, right=None)

def add_heading_apa(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run.font.italic = True
    return p

def add_body_p(doc, text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_table_caption(doc, number, title):
    p_num = doc.add_paragraph()
    p_num.paragraph_format.first_line_indent = Inches(0)
    p_num.paragraph_format.line_spacing = 2.0
    p_num.paragraph_format.space_before = Pt(12)
    p_num.paragraph_format.space_after = Pt(0)
    r_num = p_num.add_run(f"Tabla {number}")
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(12)
    r_num.font.bold = True

    p_title = doc.add_paragraph()
    p_title.paragraph_format.first_line_indent = Inches(0)
    p_title.paragraph_format.line_spacing = 2.0
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.font.italic = True

def add_placeholder_box(doc, label, description):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[ESPACIO RESERVADO PARA CAPTURA: {label}]\n{description}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.italic = True

def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # PORTADA APA 7
    for _ in range(4):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 2.0
    r_t = p_title.add_run("Sistema de Gestión de Pedidos, Inventario y Ventas:\nEntrega 2 – Sprint 2 (Compras e Inventario)")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(12)
    r_t.font.bold = True

    doc.add_paragraph()

    meta_lines = [
        "Raúl, Noé y Kevin",
        "Facultad de Ingeniería en Sistemas y Ciencias de la Computación",
        "Curso: Análisis de Sistemas II",
        "Octubre de 2026"
    ]
    for line in meta_lines:
        pm = doc.add_paragraph()
        pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pm.paragraph_format.line_spacing = 2.0
        rm = pm.add_run(line)
        rm.font.name = 'Times New Roman'
        rm.font.size = Pt(12)

    doc.add_page_break()

    # 1. META DEL SPRINT Y REGISTRO GENERAL
    add_heading_apa(doc, "1. Meta del Sprint (Sprint Goal) y Alcance de la Entrega 2", 1)
    add_body_p(doc, "El objetivo del segundo Sprint consistió en diseñar, construir y validar de extremo a extremo el proceso de abastecimiento y control de existencias del sistema de juguetería, comprendiendo el flujo continuo entre el registro de proveedores, la emisión de órdenes de compra con su respectivo detalle de productos, la confirmación de recepción de mercadería y la generación automática de entradas en el kardex de inventario. Adicionalmente, se implementó el control de niveles mínimos de existencia y el sistema automatizado de alertas por desabastecimiento, garantizando la trazabilidad transaccional y el cumplimiento estricto de las reglas de negocio definidas en el análisis del sistema.")

    # 2. PRODUCT BACKLOG ACTUALIZADO
    add_heading_apa(doc, "2. Product Backlog Actualizado", 1)
    add_body_p(doc, "En la Tabla 1 se presenta el Product Backlog general del proyecto actualizado al cierre del segundo Sprint, reflejando los elementos completados en las primeras dos iteraciones y el trabajo programado para las entregas subsecuentes.")

    add_table_caption(doc, 1, "Product Backlog General del Sistema Actualizado al Sprint 2")
    pb_data = [
        ["ID", "Historia / Requerimiento", "Módulo", "Prioridad", "Sprint", "Estado"],
        ["HU-01", "Inicio de sesión con credenciales y emisión de sesión segura", "Autenticación", "Alta", "Sprint 1", "Completado"],
        ["HU-02", "Administración de usuarios y asignación de roles", "Usuarios", "Alta", "Sprint 1", "Completado"],
        ["HU-03", "Configuración de matriz de roles y permisos (RBAC)", "Roles", "Alta", "Sprint 1", "Completado"],
        ["HU-04", "Registro, edición y consulta de cartera de clientes", "Clientes", "Media", "Sprint 1", "Completado"],
        ["HU-05", "Gestión de catálogo de productos y escalas de precios", "Productos", "Alta", "Sprint 1", "Completado"],
        ["HU-06", "Registro, edición e inactivación lógica de proveedores", "Proveedores", "Alta", "Sprint 2", "Completado"],
        ["HU-07", "Emisión de órdenes de compra con detalle de artículos", "Compras", "Alta", "Sprint 2", "Completado"],
        ["HU-08", "Recepción de mercadería e ingreso automático a inventario", "Compras / Inv.", "Alta", "Sprint 2", "Completado"],
        ["HU-09", "Consulta de existencias actuales y configuración de stock mínimo", "Inventario", "Alta", "Sprint 2", "Completado"],
        ["HU-10", "Consulta de kardex y registro justificado de ajustes/devoluciones", "Inventario", "Media", "Sprint 2", "Completado"],
        ["HU-11", "Generación y seguimiento de alertas por nivel mínimo de stock", "Inventario", "Alta", "Sprint 2", "Completado"],
        ["HU-12", "Creación de pedidos de clientes y validación de disponibilidad", "Pedidos", "Alta", "Sprint 3", "Pendiente"],
        ["HU-13", "Flujo de autorización de descuentos especiales por supervisor", "Pedidos", "Alta", "Sprint 3", "Pendiente"],
        ["HU-14", "Facturación de pedidos aprobados y salida automática de stock", "Ventas", "Alta", "Sprint 3", "Pendiente"],
        ["HU-15", "Registro de pagos parciales o totales y control de saldos", "Pagos", "Alta", "Sprint 3", "Pendiente"],
        ["HU-16", "Consulta de bitácora de auditoría inmutable por operación", "Auditoría", "Media", "Sprint 3", "Pendiente"],
        ["HU-17", "Dashboard de indicadores comerciales y exportación de reportes", "Reportes", "Media", "Sprint 4", "Pendiente"]
    ]
    t_pb = doc.add_table(rows=len(pb_data), cols=6)
    for r_i, row in enumerate(pb_data):
        for c_i, val in enumerate(row):
            t_pb.cell(r_i, c_i).text = val
    format_apa_table(t_pb)

    # 3. HISTORIAS DE USUARIO Y CRITERIOS DE ACEPTACIÓN (SPRINT 2)
    add_heading_apa(doc, "3. Historias de Usuario y Criterios de Aceptación del Sprint 2", 1)
    add_body_p(doc, "A continuación se detallan las historias de usuario seleccionadas para la segunda iteración junto con sus criterios de aceptación verificables.")

    add_table_caption(doc, 2, "Especificación de Historias de Usuario y Criterios de Aceptación del Sprint 2")
    hu_data = [
        ["ID", "Enunciado de la Historia de Usuario", "Criterios de Aceptación"],
        ["HU-06", "Como Supervisor o Administrador, quiero gestionar el catálogo de proveedores para mantener actualizados los datos fiscales y de contacto de quienes abastecen la juguetería.", "1. El sistema valida que la identificación fiscal (NIT) sea única antes de guardar.\n2. Si un proveedor tiene órdenes de compra asociadas, el sistema impide su eliminación física y aplica baja lógica mediante estado inactivo.\n3. Toda alta o modificación queda registrada en la bitácora de auditoría."],
        ["HU-07", "Como Supervisor o Administrador, quiero registrar órdenes de compra indicando el proveedor y el detalle de productos con cantidades y costos unitarios para documentar los compromisos de abastecimiento.", "1. La orden se crea inicialmente en estado REGISTRADA sin alterar aún el stock físico.\n2. El sistema exige al menos una línea de detalle con cantidad mayor a cero y costo unitario mayor a cero.\n3. No se permite registrar compras con proveedores inactivos ni incluir productos inactivos.\n4. El total de la compra se calcula automáticamente sumando los subtotales de cada renglón."],
        ["HU-08", "Como Supervisor o Administrador con permiso de recepción, quiero confirmar la recepción de una orden de compra para que el inventario de los productos adquiridos aumente de forma automática.", "1. Al confirmar la recepción, la compra cambia de estado REGISTRADA a RECIBIDA registrando fecha, hora y usuario receptor.\n2. El sistema inserta automáticamente un movimiento de tipo ENTRADA en el kardex por cada producto del detalle.\n3. El stock actual de cada producto se incrementa exactamente en la cantidad recibida.\n4. Una compra en estado RECIBIDA o ANULADA queda bloqueada contra ediciones o recepciones duplicadas."],
        ["HU-09", "Como usuario con acceso a inventario, quiero consultar las existencias actuales y configurar el stock mínimo de cada producto para prevenir quiebres de inventario.", "1. La vista muestra código SKU, nombre, categoría, stock actual, stock mínimo y cantidad faltante.\n2. El campo de stock actual no es editable de forma directa en ningún formulario.\n3. Al modificar el stock mínimo de un producto, el sistema reevalúa de inmediato si debe activarse o cerrarse una alerta."],
        ["HU-10", "Como Supervisor o Administrador, quiero consultar el kardex de movimientos y registrar ajustes o devoluciones justificadas para mantener cuadradas las existencias físicas con el sistema.", "1. El kardex expone producto, tipo de movimiento (ENTRADA, SALIDA, AJUSTE, DEVOLUCION), cantidad, stock anterior, stock posterior, documento origen, motivo y usuario.\n2. Ningún movimiento puede dejar el stock posterior en un valor negativo.\n3. Los movimientos registrados son inmutables; no pueden editarse ni borrarse."],
        ["HU-11", "Como encargado de compras e inventario, quiero visualizar alertas automáticas cuando un producto alcance o baje de su existencia mínima para gestionar su reabastecimiento oportuno.", "1. Cuando el stock actual de un producto es menor o igual a su stock mínimo, se genera automáticamente una alerta en estado ACTIVA.\n2. Cuando ingresa una compra que eleva la existencia por encima del mínimo, la alerta cambia automáticamente a estado ATENDIDA registrando la fecha de atención.\n3. El panel principal muestra un aviso con el conteo de productos en alerta activa."]
    ]
    t_hu = doc.add_table(rows=len(hu_data), cols=3)
    for r_i, row in enumerate(hu_data):
        for c_i, val in enumerate(row):
            t_hu.cell(r_i, c_i).text = val
    format_apa_table(t_hu)

    # 4. SPRINT BACKLOG Y TABLERO KANBAN DEL EQUIPO
    add_heading_apa(doc, "4. Sprint Backlog y Distribución Kanban del Equipo", 1)
    add_body_p(doc, "El trabajo técnico del Sprint 2 se dividió entre los tres integrantes del equipo de desarrollo (Raúl, Noé y Kevin), cubriendo la capa de persistencia relacional y disparadores en MySQL, los servicios transaccionales de la API REST en Node.js con Express y la interfaz de usuario en React. En la Tabla 3 se desglosa el Sprint Backlog y en la Tabla 4 se presenta la distribución del tablero Kanban al cierre de la iteración.")

    add_table_caption(doc, 3, "Sprint Backlog Técnico y Asignación de Tareas del Sprint 2")
    sb_data = [
        ["Código", "Tarea Técnica", "HU Rel.", "Responsable", "Estimación (h)", "Estado Final"],
        ["TK-01", "Diseño y ajuste de tablas proveedores, compras, compras_detalle, inventario, movimientos_inventario y alertas_inventario", "HU-06 a HU-11", "Noé", "4", "Hecho"],
        ["TK-02", "Programación de triggers trg_compras_bi/bu/au, trg_mov_bi y trg_inventario_au para entrada automática y alertas", "HU-08, HU-11", "Noé", "6", "Hecho"],
        ["TK-03", "Implementación de triggers trg_proveedores_ai/au/bd y actualización de alertas en trg_productos_au", "HU-06, HU-09", "Noé", "3", "Hecho"],
        ["TK-04", "Desarrollo del módulo backend de Proveedores (schema Zod, service con withUser, controller y routes)", "HU-06", "Raúl", "4", "Hecho"],
        ["TK-05", "Desarrollo del módulo backend de Compras (correlativo COMP-XXXX, creación maestro-detalle, recepción y anulación)", "HU-07, HU-08", "Noé", "6", "Hecho"],
        ["TK-06", "Desarrollo del módulo backend de Inventario (existencias, vista v_kardex, ajustes y consulta de alertas)", "HU-09 a HU-11", "Raúl", "5", "Hecho"],
        ["TK-07", "Construcción de vista ProveedoresPage.jsx con buscador, filtros, modal de edición y baja lógica", "HU-06", "Kevin", "4", "Hecho"],
        ["TK-08", "Construcción de vista ComprasPage.jsx con formulario maestro-detalle dinámico y modal de recepción", "HU-07, HU-08", "Kevin", "6", "Hecho"],
        ["TK-09", "Construcción de vista InventarioPage.jsx con pestañas de Existencias, Kardex y Alertas, e integración en HomePage", "HU-09 a HU-11", "Kevin", "6", "Hecho"],
        ["TK-10", "Elaboración y ejecución de suite de pruebas de integración compras_inventario.test.js en Vitest", "HU-06 a HU-11", "Raúl", "4", "Hecho"]
    ]
    t_sb = doc.add_table(rows=len(sb_data), cols=6)
    for r_i, row in enumerate(sb_data):
        for c_i, val in enumerate(row):
            t_sb.cell(r_i, c_i).text = val
    format_apa_table(t_sb)

    add_table_caption(doc, 4, "Tablero Kanban de Seguimiento de Tareas por Integrante (Raúl, Noé y Kevin)")
    kb_data = [
        ["Integrante", "Por Hacer (To Do)", "En Progreso (In Progress)", "En Revisión / Pruebas (QA)", "Completado (Done)"],
        ["Noé (Base de Datos y Arquitectura Backend)", "Ninguna al cierre del Sprint", "Ninguna al cierre del Sprint", "Verificación de bloqueo por concurrencia FOR UPDATE en trg_mov_bi", "TK-01: Esquema relacional de compras y kardex.\nTK-02: Triggers de recepción y alertas RN02/RN10.\nTK-03: Triggers de auditoría de proveedores.\nTK-05: Servicio y controlador de Compras."],
        ["Raúl (Servicios Backend, Seguridad RBAC y Testing)", "Ninguna al cierre del Sprint", "Ninguna al cierre del Sprint", "Ejecución de regresión completa (55 tests en Vitest)", "TK-04: API REST de Proveedores y validación NIT.\nTK-06: API REST de Inventario, Kardex y Ajustes.\nTK-10: Casos de prueba automatizados CP-01 a CP-07."],
        ["Kevin (Interfaces de Usuario React y Flujo UX)", "Ninguna al cierre del Sprint", "Ninguna al cierre del Sprint", "Pruebas visuales de pre-llenado desde Alertas hacia Compras", "TK-07: Pantalla ProveedoresPage.jsx.\nTK-08: Pantalla ComprasPage.jsx (Maestro-Detalle).\nTK-09: Pantalla InventarioPage.jsx y banner de alertas en inicio."]
    ]
    t_kb = doc.add_table(rows=len(kb_data), cols=5)
    for r_i, row in enumerate(kb_data):
        for c_i, val in enumerate(row):
            t_kb.cell(r_i, c_i).text = val
    format_apa_table(t_kb)

    add_placeholder_box(doc, "Tablero Kanban del Sprint 2", "Insertar captura de la herramienta de gestión (Trello / Jira / GitHub Projects) mostrando las tarjetas asignadas a Raúl, Noé y Kevin.")

    # 5. MODELO DE DATOS Y DICCIONARIO ESTRUCTURAL DEL SPRINT 2
    add_heading_apa(doc, "5. Modelo de Datos Actualizado (Módulos de Compras e Inventario)", 1)
    add_body_p(doc, "El modelo relacional se articula en torno a seis tablas principales para este incremento: proveedores, compras, compras_detalle, inventario, movimientos_inventario y alertas_inventario, apoyadas por las vistas v_catalogo, v_kardex y v_stock_bajo. La integridad referencial y las reglas de negocio se garantizan a nivel de motor mediante restricciones CHECK, claves únicas compuestas y disparadores transaccionales.")

    add_table_caption(doc, 5, "Estructura de Tablas Principales del Incremento de Compras e Inventario")
    md_data = [
        ["Tabla", "Clave Primaria", "Claves Foráneas", "Columnas Principales y Restricciones"],
        ["proveedores", "id_proveedor", "N/A", "identificacion_fiscal (UNIQUE NOT NULL), razon_social, contacto_nombre, telefono, email, direccion, activo (DEFAULT TRUE)."],
        ["compras", "id_compra", "id_proveedor -> proveedores\nid_usuario -> usuarios\nid_usuario_recepcion -> usuarios", "numero_comprobante (UNIQUE), estado ENUM('REGISTRADA', 'RECIBIDA', 'ANULADA'), total DECIMAL(12,2), observaciones, fecha_compra, fecha_recepcion."],
        ["compras_detalle", "id_compra_detalle", "id_compra -> compras\nid_producto -> productos", "cantidad INT UNSIGNED CHECK (cantidad > 0), costo_unitario DECIMAL(10,2) CHECK (costo_unitario > 0), subtotal STORED (cantidad * costo_unitario), UNIQUE(id_compra, id_producto)."],
        ["inventario", "id_producto", "id_producto -> productos", "stock_actual INT UNSIGNED NOT NULL DEFAULT 0, ultima_actualizacion TIMESTAMP. Protegida contra edición manual mediante variable de sesión @inv_trigger."],
        ["movimientos_inventario", "id_movimiento", "id_producto -> productos\nid_usuario -> usuarios", "tipo_movimiento ENUM('ENTRADA', 'SALIDA', 'AJUSTE', 'DEVOLUCION'), cantidad, stock_anterior, stock_posterior, documento_tipo, documento_id, motivo, fecha_hora."],
        ["alertas_inventario", "id_alerta", "id_producto -> productos", "stock_al_generar INT, stock_minimo INT UNSIGNED, estado ENUM('ACTIVA', 'ATENDIDA'), fecha_generacion, fecha_atencion."]
    ]
    t_md = doc.add_table(rows=len(md_data), cols=4)
    for r_i, row in enumerate(md_data):
        for c_i, val in enumerate(row):
            t_md.cell(r_i, c_i).text = val
    format_apa_table(t_md)

    # 6. DIAGRAMAS ACTUALIZADOS
    add_heading_apa(doc, "6. Diagramas Actualizados del Sistema", 1)
    add_body_p(doc, "En esta sección se incorporan los diagramas actualizados correspondientes al segundo incremento funcional: el diagrama de casos de uso del subsistema de abastecimiento e inventario, el diagrama de secuencia del proceso de recepción de compra con actualización de kardex y el modelo entidad-relación actualizado.")

    add_placeholder_box(doc, "Diagrama de Casos de Uso - Sprint 2 (Compras e Inventario)", "Insertar imagen del diagrama de casos de uso mostrando los actores Administrador, Supervisor y Operador interactuando con Gestionar Proveedores, Registrar Compra, Recibir Mercadería, Consultar Kardex, Ajustar Inventario y Gestionar Alertas.")
    add_placeholder_box(doc, "Diagrama de Secuencia - Proceso de Recepción y Entrada de Inventario", "Insertar imagen del diagrama de secuencia: Usuario -> ComprasPage -> ComprasController -> ComprasService (withUser) -> MySQL (trg_compras_au -> trg_mov_bi -> trg_inventario_au).")
    add_placeholder_box(doc, "Modelo Entidad-Relación Actualizado - Sprint 2", "Insertar imagen del diagrama Entidad-Relación destacando las relaciones entre proveedores, compras, compras_detalle, productos, inventario, movimientos_inventario y alertas_inventario.")

    # 7. REGLAS DE NEGOCIO IMPLEMENTADAS
    add_heading_apa(doc, "7. Reglas de Negocio Implementadas en el Sprint 2", 1)
    add_body_p(doc, "Aunque los lineamientos de la segunda entrega establecen un mínimo de cinco reglas de negocio, el equipo implementó siete reglas verificables tanto en la capa de servicios como en los disparadores de la base de datos para blindar la operación frente a inconsistencias.")

    add_table_caption(doc, 6, "Catálogo de Reglas de Negocio Implementadas y Mecanismo de Control")
    rn_data = [
        ["Código", "Regla de Negocio", "Implementación Técnica (BD / Backend)", "Resultado ante Incumplimiento"],
        ["RN-01", "Prohibición de existencias negativas en inventario", "Trigger trg_mov_bi evalúa v_nuevo = v_stock + v_delta antes de aplicar cualquier salida o ajuste negativo.", "Transacción abortada con SQLSTATE 45000 (HTTP 409): Inventario insuficiente."],
        ["RN-02", "Aumento automático de inventario al recibir una compra", "Trigger trg_compras_au detecta el paso de REGISTRADA a RECIBIDA e inserta líneas ENTRADA en movimientos_inventario.", "Incremento automático e inmediato de inventario.stock_actual y registro en kardex."],
        ["RN-04", "Prohibición de eliminación física de productos o proveedores con historial", "Triggers trg_productos_bd y trg_proveedores_bd verifican existencia en compras y movimientos.", "Se bloquea el DELETE físico y se exige el uso de baja lógica (activo = 0)."],
        ["RN-10", "Generación y cierre automático de alertas por stock mínimo", "Triggers trg_inventario_au y trg_productos_au comparan stock_actual <= stock_minimo tras cada cambio.", "Se crea registro ACTIVA en alertas_inventario o pasa a ATENDIDA al superar el umbral."],
        ["RN-11", "Inmutabilidad manual del stock de inventario", "Trigger trg_inventario_bu verifica la variable interna @inv_trigger = 1 establecida solo desde trg_mov_bi.", "Cualquier UPDATE directo sobre la tabla inventario es rechazado por el motor SQL."],
        ["RN-12", "Inmutabilidad de órdenes de compra finalizadas", "Trigger trg_compras_bu y validaciones en ComprasService.actualizar / recibir.", "Una compra en estado RECIBIDA o ANULADA no admite cambios ni dobles recepciones."],
        ["RN-13", "Restricción de abastecimiento solo con proveedores y productos activos", "Validación en ComprasService.crear y trigger trg_cdet_bi sobre el campo activo.", "Se rechaza el registro de la compra si el proveedor o algún artículo está inactivo."]
    ]
    t_rn = doc.add_table(rows=len(rn_data), cols=4)
    for r_i, row in enumerate(rn_data):
        for c_i, val in enumerate(row):
            t_rn.cell(r_i, c_i).text = val
    format_apa_table(t_rn)

    # 8. CASOS DE PRUEBA EJECUTADOS
    add_heading_apa(doc, "8. Matriz y Ejecución de Casos de Prueba del Sprint 2", 1)
    add_body_p(doc, "Se diseñaron y ejecutaron siete casos de prueba automatizados dentro del archivo backend/tests/compras_inventario.test.js utilizando Vitest y Supertest contra la base de datos de pruebas jugueteria_test. En la Tabla 7 se documentan los escenarios evaluados.")

    add_table_caption(doc, 7, "Casos de Prueba Automatizados Ejecutados en el Sprint 2")
    cp_data = [
        ["ID Caso", "Historia / Regla", "Descripción y Datos de Entrada", "Resultado Esperado", "Resultado Obtenido", "Estado"],
        ["CP-01", "HU-06", "Crear proveedor con NIT '9988776-5' e intentar registrar un segundo proveedor con el mismo NIT.", "Primer registro retorna HTTP 201; el segundo retorna HTTP 409 por NIT duplicado.", "HTTP 201 en alta inicial y HTTP 409 (DUPLICATE_NIT) en el duplicado.", "Aprobado"],
        ["CP-02", "HU-07", "Registrar orden de compra en estado REGISTRADA con 10 unidades de producto y consultar stock actual.", "La compra se guarda con total calculado, pero stock_actual permanece sin cambios hasta la recepción.", "Compra creada en estado REGISTRADA; el stock se mantuvo intacto.", "Aprobado"],
        ["CP-03", "HU-08 / RN-02", "Ejecutar PATCH /api/compras/:id/recibir sobre la orden registrada y consultar inventario y kardex.", "Estado pasa a RECIBIDA, stock_actual aumenta en +10 y se crea movimiento ENTRADA.", "Stock incrementado exactamente en 10 unidades y movimiento ENTRADA registrado.", "Aprobado"],
        ["CP-04", "HU-08 / RN-12", "Intentar modificar o volver a ejecutar la recepción sobre una compra que ya se encuentra en estado RECIBIDA.", "El sistema bloquea la operación con código HTTP 409 indicando que la compra es inmutable.", "Operación rechazada con HTTP 409 (BUSINESS_RULE_VIOLATION).", "Aprobado"],
        ["CP-05", "HU-11 / RN-10", "Subir el stock_minimo por encima del stock_actual para provocar alerta y luego recibir compra de abastecimiento.", "Se genera alerta ACTIVA al quedar bajo el mínimo y cambia a ATENDIDA tras recibir la compra.", "Transición automática ACTIVA -> ATENDIDA verificada en base de datos.", "Aprobado"],
        ["CP-06", "HU-10 / RN-01", "Intentar registrar un ajuste negativo mayor a la existencia actual (-9999) y probar acceso con rol Operador.", "El ajuste excesivo es rechazado (HTTP 409) y el rol Operador recibe HTTP 403 Forbidden.", "Ambos controles bloquearon la operación según lo especificado.", "Aprobado"],
        ["CP-07", "HU-06 / RN-04", "Intentar eliminar físicamente (DELETE SQL) un proveedor que ya posee órdenes de compra asociadas.", "El trigger trg_proveedores_bd aborta el borrado físico exigiendo pasar a estado inactivo.", "Excepción SQLSTATE 45000 capturada correctamente.", "Aprobado"]
    ]
    t_cp = doc.add_table(rows=len(cp_data), cols=6)
    for r_i, row in enumerate(cp_data):
        for c_i, val in enumerate(row):
            t_cp.cell(r_i, c_i).text = val
    format_apa_table(t_cp)

    # 9. EVIDENCIA DE DAILY SCRUM
    add_heading_apa(doc, "9. Evidencia de Sesiones de Daily Scrum", 1)
    add_body_p(doc, "Durante el desarrollo del Sprint 2, el equipo realizó reuniones breves de sincronización para inspeccionar el avance hacia la meta del Sprint y resolver impedimentos técnicos de integración entre base de datos, backend y frontend.")

    add_table_caption(doc, 8, "Bitácora de Sincronización Daily Scrum del Equipo (Raúl, Noé y Kevin)")
    ds_data = [
        ["Sesión / Fecha", "Integrante", "¿Qué hice ayer?", "¿Qué haré hoy?", "¿Existe algún impedimento?"],
        ["Sesión 1\n(Inicio de Sprint)", "Noé", "Revisión del esquema SQL base del Sprint 1.", "Programar los triggers de compras, recepción y movimientos de kardex en MySQL.", "Ninguno."],
        ["Sesión 1\n(Inicio de Sprint)", "Raúl", "Definición de contratos de rutas y permisos para Sprint 2.", "Desarrollar schemas Zod y servicios de Proveedores.", "Ninguno."],
        ["Sesión 1\n(Inicio de Sprint)", "Kevin", "Diseño de estructura de pantallas para Proveedores y Compras.", "Construir el componente ProveedoresPage.jsx conectado a la API.", "Esperando definición final de respuesta del endpoint de proveedores."],
        ["Sesión 2\n(Mitad de Sprint)", "Noé", "Triggers de compras y kardex listos; detecté que faltaban triggers de auditoría en proveedores.", "Agregar trg_proveedores_ai/au/bd y desarrollar el servicio de Compras con correlativo COMP-XXXX.", "Al cambiar stock_minimo en productos no se disparaba la alerta si no cambiaba inventario; lo resolveré en trg_productos_au."],
        ["Sesión 2\n(Mitad de Sprint)", "Raúl", "Módulo de Proveedores terminado y probado.", "Implementar el módulo de Inventario (existencias, kardex, ajustes y alertas) y mapear errores SQL 45000 en errorHandler.", "Ninguno."],
        ["Sesión 2\n(Mitad de Sprint)", "Kevin", "ProveedoresPage.jsx completada.", "Desarrollar ComprasPage.jsx con tabla dinámica de detalles y cálculo de subtotales.", "Necesitaba permitir que AuthContext.can() evaluara arreglos de permisos; ya lo ajusté."],
        ["Sesión 3\n(Cierre de Sprint)", "Noé", "Servicio de compras y recepción inmediata terminados y scripts reset.ps1 actualizados.", "Apoyar en las pruebas de integración de extremos con la base de datos de test.", "Ninguno."],
        ["Sesión 3\n(Cierre de Sprint)", "Raúl", "Endpoints de inventario y alertas terminados.", "Escribir y ejecutar los 7 tests automatizados en compras_inventario.test.js.", "Ninguno; los 55 tests del proyecto pasan en verde."],
        ["Sesión 3\n(Cierre de Sprint)", "Kevin", "ComprasPage.jsx e InventarioPage.jsx con sus 3 pestañas terminadas.", "Conectar el botón Abastecer desde las alertas hacia el modal de compras y verificar npm run build.", "Ninguno."]
    ]
    t_ds = doc.add_table(rows=len(ds_data), cols=5)
    for r_i, row in enumerate(ds_data):
        for c_i, val in enumerate(row):
            t_ds.cell(r_i, c_i).text = val
    format_apa_table(t_ds)

    add_placeholder_box(doc, "Evidencia de Reunión Daily Scrum", "Insertar captura de pantalla de la videollamada o hilo de comunicación del equipo (Raúl, Noé y Kevin) durante el seguimiento del Sprint 2.")

    # 10. SPRINT REVIEW Y EVIDENCIA FUNCIONAL
    add_heading_apa(doc, "10. Revisión del Sprint (Sprint Review) y Evidencia Funcional", 1)
    add_body_p(doc, "Al concluir la iteración se llevó a cabo la revisión del incremento funcional verificando el cumplimiento de la Definición de Terminado (Definition of Done): código integrado en la rama principal, compilación de producción sin errores en React (Vite) y aprobación del 100% de las pruebas automatizadas (55 pruebas aprobadas en 6 suites). Todas las historias comprometidas (HU-06 a HU-11) fueron aceptadas.")

    add_placeholder_box(doc, "Pantalla de Gestión de Proveedores (HU-06)", "Insertar captura de ProveedoresPage mostrando el listado de proveedores, el buscador y el modal de creación/edición.")
    add_placeholder_box(doc, "Registro de Orden de Compra y Detalle (HU-07)", "Insertar captura del modal Nueva Compra en ComprasPage mostrando el proveedor seleccionado, las líneas de productos y el cálculo del total.")
    add_placeholder_box(doc, "Recepción de Compra y Estado RECIBIDA (HU-08)", "Insertar captura del modal de Detalle y Recepción de Compra confirmando el ingreso de mercadería al almacén.")
    add_placeholder_box(doc, "Centro de Control de Inventario y Kardex de Movimientos (HU-09 y HU-10)", "Insertar captura de InventarioPage mostrando el aumento automático de stock y el movimiento de tipo ENTRADA en el Kardex.")
    add_placeholder_box(doc, "Alertas de Stock Mínimo Activas y Atendidas (HU-11)", "Insertar captura de la pestaña Alertas de Stock y el banner de alerta en el Dashboard principal.")

    # 11. SPRINT RETROSPECTIVE
    add_heading_apa(doc, "11. Retrospectiva del Sprint (Sprint Retrospective)", 1)
    add_body_p(doc, "El equipo conformado por Raúl, Noé y Kevin realizó la sesión de retrospectiva analizando el desempeño técnico y colaborativo durante el segundo Sprint, cuyas conclusiones se sintetizan en la Tabla 9.")

    add_table_caption(doc, 9, "Matriz de Retrospectiva del Equipo para el Sprint 2")
    ret_data = [
        ["Dimensión Analizada", "Observaciones del Equipo", "Plan de Acción para el Sprint 3 (Pedidos y Ventas)"],
        ["¿Qué salió bien?", "1. Delegar la lógica crítica de actualización de inventario y generación de alertas a triggers transaccionales de MySQL evitó condiciones de carrera y garantizó que ningún endpoint pudiera alterar el stock manualmente.\n2. La reutilización de los componentes base creados en el Sprint 1 (DataTable, Modal, PageCard) aceleró el desarrollo de las vistas de Proveedores, Compras e Inventario.\n3. La suite de pruebas automatizadas permitió detectar y corregir problemas de permisos y códigos de estado antes de integrar el frontend.", "Mantener el patrón de diseño basado en triggers y transacciones envueltas en withUser para el flujo de Pedidos, Autorización de Descuentos, Ventas y Pagos del Sprint 3."],
        ["¿Qué salió mal o presentó dificultades?", "1. Inicialmente, al cambiar el stock_minimo desde la tabla de productos no se regeneraban las alertas hasta que ocurría un nuevo movimiento en el kardex.\n2. La función de verificación de permisos en el frontend solo admitía un permiso simple y no una lista de permisos alternativos para rutas compartidas entre roles.", "Revisar desde la fase de diseño qué eventos secundarios (como cambios de parámetros o anulaciones) deben disparar reevaluaciones de estado en cascada."],
        ["¿Qué debemos mejorar?", "1. Sincronizar desde el primer día los formatos de respuesta de los endpoints antes de construir los formularios complejos maestro-detalle.\n2. Automatizar la captura de evidencias visuales conforme se cierra cada historia de usuario para no acumularlas al final de la entrega.", "Documentar los contratos JSON de los endpoints de Pedidos, Ventas y Pagos al inicio del Sprint 3 y tomar las capturas inmediatamente después de aprobar cada caso de prueba."]
    ]
    t_ret = doc.add_table(rows=len(ret_data), cols=3)
    for r_i, row in enumerate(ret_data):
        for c_i, val in enumerate(row):
            t_ret.cell(r_i, c_i).text = val
    format_apa_table(t_ret)

    os.makedirs("docs/sprint-2", exist_ok=True)
    out_path = "docs/sprint-2/Entrega_2_Sprint_2_APA7.docx"
    doc.save(out_path)
    print(f"Documento generado exitosamente en: {out_path}")

if __name__ == "__main__":
    build_document()
